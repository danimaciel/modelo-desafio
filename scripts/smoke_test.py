import os
import sys
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ.setdefault('HF_HOME', str(ROOT / '.cache' / 'huggingface'))
os.chdir(ROOT)

import openpyxl
from reportlab.pdfgen import canvas
from reportlab.platypus import Table, TableStyle
from reportlab.lib import colors
from streamlit.testing.v1 import AppTest
from desafios.catalog import load_catalog
from desafios.files import load_xlsx, read_pdf
from desafios.semantic import SemanticEngine


def click(app, label):
    next(button for button in app.button if button.label == label).click().run(timeout=90)
    assert not app.exception, str(app.exception)
    assert not app.error, [e.value for e in app.error]


catalog, version = load_catalog()
engine = SemanticEngine(catalog, version)
for row in [catalog[0], catalog[4], catalog[40], catalog[-1]]:
    result = engine.rank(row['desafio'])[0]
    assert result['id_desafio'] == row['id_desafio'], (row['id_desafio'], result['id_desafio'])
print('SBERT CPU: quatro autoconsultas e metadados OK')

workbook = openpyxl.Workbook()
sheet = workbook.active
sheet.title = 'Publicacoes'
sheet.append(['Titulo', 'Resumo', 'Autor', 'Formula'])
sheet.append(['Tema agroambiental', catalog[0]['desafio'], 'Teste', '=1+2'])
sheet.append(['', '', 'Linha sem texto', None])
workbook.create_sheet('Preservada')['A1'] = 'original'
upload = BytesIO()
workbook.save(upload)
upload.name = 'publicacoes_teste.xlsx'
with patch('streamlit.file_uploader', return_value=upload):
    app = AppTest.from_file(str(ROOT / 'app.py')).run(timeout=30)
    assert not app.error, [e.value for e in app.error]
    click(app, 'Analisar linhas')
    result = app.session_state['result']
    output = load_xlsx(result['xlsx'])
    assert output['Publicacoes']['E2'].value.startswith('D001')
    assert 'Sem texto' in output['Publicacoes']['E3'].value
    assert output['Publicacoes']['D2'].value == '=1+2'
    assert output['Preservada']['A1'].value == 'original'
    app.slider[0].set_value(0.5).run()
    assert not app.success, 'Stale results were shown after a parameter change'
print('UI XLSX: selecao, duas linhas, copia, formulas, outras abas e invalidacao OK')

pdf = BytesIO()
c = canvas.Canvas(pdf)
c.drawString(40, 790, 'Projeto: sistemas produtivos sustentaveis e servicos ecossistemicos')
text = c.beginText(40, 760)
for line in ['Objetivo: ampliar a capacidade adaptativa dos sistemas produtivos.',
             'Promover controle de emissoes de GEE e servicos ecossistemicos.',
             'Agregar valor a produtos e servicos de sistemas agricolas tradicionais.']:
    text.textLine(line)
c.drawText(text)
c.showPage()
c.save()
pdf.name = 'projeto_teste.pdf'
with patch('streamlit.file_uploader', return_value=pdf):
    app = AppTest.from_file(str(ROOT / 'app.py')).run(timeout=30)
    app.radio[0].set_value('Projeto em PDF').run()
    click(app, 'Extrair texto do projeto')
    click(app, 'Analisar projeto')
    results = app.session_state['result']['details']['resultados']
    assert 1 <= len(results) <= 3
    assert results[0]['evidence'][0]['page'] == 1
print('UI PDF projeto: extracao, ranking e evidencia paginada OK')

table_pdf = BytesIO()
c = canvas.Canvas(table_pdf)
table = Table([['Titulo', 'Descricao'], ['Tecnologia', 'Manejo sustentavel do solo'],
               ['Publicacao', 'Producao de leite com qualidade']], colWidths=[120, 300])
table.setStyle(TableStyle([('GRID', (0, 0), (-1, -1), 1, colors.black)]))
table.wrapOn(c, 500, 500)
table.drawOn(c, 40, 600)
c.save()
table_pdf.name = 'tabela_teste.pdf'
with patch('streamlit.file_uploader', return_value=table_pdf):
    app = AppTest.from_file(str(ROOT / 'app.py')).run(timeout=30)
    click(app, 'Extrair tabelas')
    app.checkbox[0].check().run()
    click(app, 'Analisar tabela')
    output = load_xlsx(app.session_state['result']['xlsx'])
    assert output.active.max_row == 3
    assert output.active.max_column == 3
print('UI PDF tabela: extracao, confirmacao, analise e exportacao OK')

blank = BytesIO()
c = canvas.Canvas(blank)
c.showPage()
c.save()
try:
    read_pdf(blank.getvalue())
    raise AssertionError('Blank PDF was accepted')
except ValueError as e:
    assert 'OCR' in str(e)
print('PDF sem texto: bloqueio explicito OK')
