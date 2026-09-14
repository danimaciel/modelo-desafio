from copy import copy
from io import BytesIO, StringIO
import csv
from zipfile import ZipFile, BadZipFile
import unicodedata

import openpyxl
import pdfplumber

MAX_BYTES = 20 * 1024 * 1024
MAX_ROWS = 2000
MAX_COLUMNS = 100
MAX_PAGES = 150


def prepare_table(raw, extension, delimiter=';', encoding='utf-8-sig'):
    """Converte XLS/CSV em uma representação XLSX para o fluxo de análise."""
    validate_bytes(raw)
    if extension == '.xlsx':
        return raw
    workbook = openpyxl.Workbook()
    if extension == '.csv':
        try:
            text = raw.decode(encoding)
        except UnicodeError as exc:
            raise ValueError('Não foi possível ler os caracteres. Selecione outra codificação do CSV.') from exc
        sheet = workbook.active
        sheet.title = 'Dados'
        try:
            for number, row in enumerate(csv.reader(StringIO(text, newline=''), delimiter=delimiter, strict=True), 1):
                if number > MAX_ROWS + 100 or len(row) > MAX_COLUMNS:
                    raise ValueError('CSV excede 2.100 linhas ou 100 colunas.')
                if any(len(value) > 32767 for value in row):
                    raise ValueError('Uma célula excede 32.767 caracteres. Divida seu conteúdo antes de analisar.')
                sheet.append(row)
                for cell in sheet[number]:
                    cell.data_type = 's'
        except csv.Error as exc:
            raise ValueError('CSV inválido. Confira o separador e as aspas do arquivo.') from exc
    elif extension == '.xls':
        import xlrd
        try:
            source = xlrd.open_workbook(file_contents=raw, on_demand=True)
        except xlrd.XLRDError as exc:
            raise ValueError('XLS inválido ou protegido. Salve uma cópia sem senha no Excel.') from exc
        try:
            workbook.remove(workbook.active)
            for original in source.sheets():
                if original.nrows > MAX_ROWS + 100 or original.ncols > MAX_COLUMNS:
                    raise ValueError('XLS excede 2.100 linhas ou 100 colunas por aba.')
                sheet = workbook.create_sheet(original.name)
                for row in original.get_rows():
                    values = []
                    for cell in row:
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, source.datemode)
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            value = xlrd.error_text_from_code.get(value, '#ERRO')
                        values.append(value)
                    sheet.append(values)
                    for cell in sheet[sheet.max_row]:
                        if isinstance(cell.value, str):
                            cell.data_type = 's'
        finally:
            source.release_resources()
    else:
        raise ValueError('Formato não suportado. Use XLSX, XLS ou CSV.')
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_csv(raw_xlsx, sheet_name, delimiter):
    sheet = load_xlsx(raw_xlsx)[sheet_name]
    output = StringIO(newline='')
    writer = csv.writer(output, delimiter=delimiter)
    for row in sheet.iter_rows(values_only=True):
        # Evita execução de fórmulas ao abrir texto importado no Excel.
        writer.writerow(["'" + v if isinstance(v, str) and v.lstrip().startswith(('=', '+', '-', '@'))
                         else ('' if v is None else v) for v in row])
    return output.getvalue().encode('utf-8-sig')


def validate_bytes(raw):
    if not raw or len(raw) > MAX_BYTES:
        raise ValueError('Envie um arquivo não vazio de até 20 MB.')


def load_xlsx(raw, data_only=False):
    validate_bytes(raw)
    try:
        with ZipFile(BytesIO(raw)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Planilha descompactada excede 100 MB.')
            names = archive.namelist()
            if any('vbaProject' in n or n.startswith('xl/embeddings/') for n in names):
                raise ValueError('Remova macros e objetos incorporados antes de enviar.')
        workbook = openpyxl.load_workbook(BytesIO(raw), data_only=data_only, keep_links=True)
    except (BadZipFile, KeyError) as exc:
        raise ValueError('O arquivo não é uma planilha XLSX válida.') from exc
    for sheet in workbook:
        if sheet.max_row > MAX_ROWS + 100 or sheet.max_column > MAX_COLUMNS:
            raise ValueError('Limite: 2.100 linhas e 100 colunas por aba, incluindo formatação.')
    return workbook


def headers(sheet, row):
    return [f'{openpyxl.utils.get_column_letter(i)} — {sheet.cell(row, i).value or "Sem título"}'
            for i in range(1, sheet.max_column + 1)]


def suggested_columns(labels):
    words = ('titulo', 'resumo', 'descricao', 'palavra', 'objetivo', 'resultado',
             'compromisso', 'solucao', 'tecnologia', 'publicacao')
    return [i for i, label in enumerate(labels) if any(w in unicodedata.normalize(
        'NFKD', label.lower()).encode('ascii', 'ignore').decode() for w in words)]


def row_text(sheet, row, columns):
    return '\n'.join(str(sheet.cell(row, col).value).strip() for col in columns
                     if sheet.cell(row, col).value is not None
                     and str(sheet.cell(row, col).value).strip())


def export_xlsx(raw, sheet_name, header_row, results):
    """Edita a cópia em memória; preserva fórmulas e outras abas com openpyxl."""
    workbook = load_xlsx(raw)
    sheet = workbook[sheet_name]
    column = sheet.max_column + 1
    if column > MAX_COLUMNS:
        raise ValueError('Não há espaço para a coluna de resultado no limite do protótipo.')
    existing = {str(c.value) for c in sheet[header_row]}
    title = 'Enquadramento sugerido'
    suffix = 2
    while title in existing:
        title = f'Enquadramento sugerido ({suffix})'
        suffix += 1
    sheet.cell(header_row, column, title)
    sheet.cell(header_row, column).font = copy(sheet.cell(header_row, 1).font)
    sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 65
    for row, value in results.items():
        cell = sheet.cell(row, column)
        cell.value = value
        cell.data_type = 's'
        cell.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical='top')
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def read_pdf(raw, tables=False, strategy='lines'):
    validate_bytes(raw)
    pages, found, empty = [], [], []
    with pdfplumber.open(BytesIO(raw)) as pdf:
        if len(pdf.pages) > MAX_PAGES:
            raise ValueError('Limite de 150 páginas por PDF.')
        for number, page in enumerate(pdf.pages, 1):
            content = page.extract_text() or ''
            if not content.strip():
                empty.append(number)
            pages.append((number, content))
            if tables:
                settings = {'vertical_strategy': strategy, 'horizontal_strategy': strategy}
                for table_number, table in enumerate(page.extract_tables(settings), 1):
                    if table:
                        found.append({'page': number, 'table': table_number, 'rows': table})
            page.close()
    if not any(text.strip() for _, text in pages):
        raise ValueError('PDF sem texto extraível. Faça OCR e envie novamente; OCR não está incluído nesta versão.')
    return pages, found, empty


def export_pdf_table(rows, results):
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Tabela analisada'
    for row in rows:
        sheet.append([str(v) if v is not None else '' for v in row])
    # Conteúdo extraído é texto, nunca fórmula executável.
    for row in sheet:
        for cell in row:
            cell.data_type = 's'
    output = BytesIO()
    workbook.save(output)
    return export_xlsx(output.getvalue(), sheet.title, 1, results)
